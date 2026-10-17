"""Exportação para uma planilha Excel nova (não altera nenhuma planilha existente).

Aba "Controle": as colunas de destino, uma linha por item; valores pendentes saem
vazios e destacados. Aba "Rastreabilidade": origem, evidência e pendência de cada valor.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from automatizador.dominio.modelo import Linha

DESTAQUE_PENDENTE = PatternFill("solid", start_color="FFF2CC", end_color="FFF2CC")
COLUNAS_RASTREABILIDADE = [
    "DOCUMENTO",
    "LINHA",
    "COLUNA",
    "VALOR SUGERIDO",
    "ORIGEM",
    "PÁGINA",
    "TRECHO",
    "CONFIANÇA",
    "PENDÊNCIA",
]


class ExportadorXlsx:
    def exportar(self, linhas: list[tuple[str, Linha]], colunas: list[str], destino: Path) -> None:
        destino = Path(destino)
        destino.parent.mkdir(parents=True, exist_ok=True)
        livro = Workbook()

        controle = livro.active
        controle.title = "Controle"
        controle.append(colunas)
        for _, linha in linhas:
            _adicionar_como_texto(controle, [linha.valores[c].valor_para_exportar for c in colunas])
        for numero, (_, linha) in enumerate(linhas, start=2):
            for indice, coluna in enumerate(colunas, start=1):
                if not linha.valores[coluna].confirmado:
                    controle.cell(row=numero, column=indice).fill = DESTAQUE_PENDENTE
        _formatar_como_tabela(controle, "Controle", len(colunas), len(linhas))

        rastreio = livro.create_sheet("Rastreabilidade")
        rastreio.append(COLUNAS_RASTREABILIDADE)
        for documento, linha in linhas:
            for coluna in colunas:
                v = linha.valores[coluna]
                _adicionar_como_texto(
                    rastreio,
                    [documento, linha.ordem, coluna, v.valor, str(v.origem), v.pagina, v.trecho, v.confianca, v.pendencia],
                )
        _formatar_como_tabela(
            rastreio, "Rastreabilidade", len(COLUNAS_RASTREABILIDADE), len(linhas) * len(colunas)
        )

        livro.save(destino)


def _adicionar_como_texto(aba, valores: list) -> None:
    """Acrescenta uma linha garantindo que nenhum texto vindo do documento vire fórmula.

    Um valor como "=HIPERLINK(...)" num PDF recebido seria executado pelo Excel se gravado como fórmula.
    """
    aba.append(valores)
    for celula in aba[aba.max_row]:
        if isinstance(celula.value, str) and celula.data_type == "f":
            celula.data_type = "s"


def _formatar_como_tabela(aba, nome: str, total_colunas: int, total_linhas: int) -> None:
    for celula in aba[1]:
        celula.font = Font(bold=True)
    for indice in range(1, total_colunas + 1):
        aba.column_dimensions[get_column_letter(indice)].width = 18
    aba.freeze_panes = "A2"
    if total_linhas:
        intervalo = f"A1:{get_column_letter(total_colunas)}{total_linhas + 1}"
        tabela = Table(displayName=nome, ref=intervalo)
        tabela.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        aba.add_table(tabela)
