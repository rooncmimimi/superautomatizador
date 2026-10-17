"""Testes no ponto principal: o caso de uso de processar documentos.

A IA é substituída por respostas gravadas; leitor de PDF, banco e exportador são os reais.
"""

import json

import pytest

from openpyxl import load_workbook

from automatizador.adaptadores.exportador_xlsx import ExportadorXlsx
from automatizador.adaptadores.ia_simulada import IASimulada
from automatizador.adaptadores.leitor_pdf import LeitorPdf
from automatizador.adaptadores.repositorio_sqlite import RepositorioSqlite
from automatizador.aplicacao.processar_documento import ProcessarDocumentos
from automatizador.dominio.colunas import COLUNAS_CONTROLE


def _caso_de_uso(tmp_path, pasta_respostas):
    repositorio = RepositorioSqlite(tmp_path / "automatizador.db")
    return ProcessarDocumentos(
        leitor=LeitorPdf(),
        assistente=IASimulada(pasta_respostas),
        repositorio=repositorio,
        exportador=ExportadorXlsx(),
    ), repositorio


def _linhas_da_planilha(caminho):
    aba = load_workbook(caminho)["Controle"]
    linhas = list(aba.iter_rows(values_only=True))
    return linhas[0], linhas[1:]


def test_cte_ficticio_vira_uma_linha_com_as_18_colunas(tmp_path, cte_ficticio):
    caso, _ = _caso_de_uso(tmp_path, cte_ficticio.pdf.parent)
    saida = tmp_path / "saida.xlsx"

    caso.executar([cte_ficticio.pdf], saida)

    cabecalho, linhas = _linhas_da_planilha(saida)
    assert list(cabecalho) == COLUNAS_CONTROLE
    assert len(linhas) == 1
    linha = dict(zip(cabecalho, linhas[0]))
    assert linha["CT-e"] == "1234"
    assert linha["FORNECEDOR"] == "Transportadora Exemplo Ltda"
    assert linha["VALOR TOTAL"] == "1.850,40"
    assert linha["DATA DA NOTA"] == "03/10/2026"


def test_valores_da_planilha_conferem_com_o_gabarito(tmp_path, cte_ficticio):
    caso, _ = _caso_de_uso(tmp_path, cte_ficticio.pdf.parent)
    saida = tmp_path / "saida.xlsx"

    caso.executar([cte_ficticio.pdf], saida)

    gabarito = json.loads(cte_ficticio.gabarito.read_text(encoding="utf-8"))
    cabecalho, linhas = _linhas_da_planilha(saida)
    for esperada, obtida in zip(gabarito["linhas"], linhas):
        obtida = dict(zip(cabecalho, obtida))
        for coluna in COLUNAS_CONTROLE:
            assert obtida[coluna] == esperada.get(coluna), coluna


def test_cada_valor_tem_origem_e_evidencia_no_banco(tmp_path, cte_ficticio):
    caso, repositorio = _caso_de_uso(tmp_path, cte_ficticio.pdf.parent)

    resultado = caso.executar([cte_ficticio.pdf], tmp_path / "saida.xlsx")

    valores = repositorio.valores_do_documento(resultado.documentos[0].id)
    preenchidos = [v for v in valores if v.valor is not None and v.coluna != "TIPO DE NOTA"]
    assert preenchidos
    for valor in preenchidos:
        assert valor.origem == "sugerido_ia"
        assert valor.trecho
        assert valor.pagina == 1
    assert repositorio.execucoes()[-1].resultado == "concluida"


def test_valor_sem_evidencia_no_documento_fica_pendente_e_vazio(tmp_path, cte_ficticio):
    resposta = json.loads(cte_ficticio.resposta_simulada.read_text(encoding="utf-8"))
    resposta["linhas"][0]["campos"]["PEDIDO"] = {
        "valor": "4500099999",
        "trecho": "Pedido: 4500099999",
        "pagina": 1,
        "confianca": 0.9,
    }
    cte_ficticio.resposta_simulada.write_text(json.dumps(resposta), encoding="utf-8")
    caso, repositorio = _caso_de_uso(tmp_path, cte_ficticio.pdf.parent)
    saida = tmp_path / "saida.xlsx"

    resultado = caso.executar([cte_ficticio.pdf], saida)

    cabecalho, linhas = _linhas_da_planilha(saida)
    assert dict(zip(cabecalho, linhas[0]))["PEDIDO"] is None
    pedido = next(
        v for v in repositorio.valores_do_documento(resultado.documentos[0].id) if v.coluna == "PEDIDO"
    )
    assert pedido.pendencia == "trecho não encontrado no documento"


def test_valor_que_nao_aparece_no_trecho_fica_pendente(tmp_path, cte_ficticio):
    resposta = json.loads(cte_ficticio.resposta_simulada.read_text(encoding="utf-8"))
    campo = resposta["linhas"][0]["campos"]["VALOR TOTAL"]
    campo["valor"] = "9.999,99"
    cte_ficticio.resposta_simulada.write_text(json.dumps(resposta), encoding="utf-8")
    caso, repositorio = _caso_de_uso(tmp_path, cte_ficticio.pdf.parent)

    resultado = caso.executar([cte_ficticio.pdf], tmp_path / "saida.xlsx")

    valor = next(
        v for v in repositorio.valores_do_documento(resultado.documentos[0].id) if v.coluna == "VALOR TOTAL"
    )
    assert valor.pendencia == "valor não aparece no trecho citado"


def test_campo_ausente_fica_registrado_como_ausente(tmp_path, cte_ficticio):
    caso, repositorio = _caso_de_uso(tmp_path, cte_ficticio.pdf.parent)

    resultado = caso.executar([cte_ficticio.pdf], tmp_path / "saida.xlsx")

    codigo = next(
        v
        for v in repositorio.valores_do_documento(resultado.documentos[0].id)
        if v.coluna == "CÓDIGO FORNECEDOR"
    )
    assert codigo.valor is None
    assert codigo.origem == "ausente"
    assert codigo.pendencia == "não encontrado no documento"


def test_resposta_fora_do_contrato_falha_a_execucao_sem_gerar_planilha(tmp_path, cte_ficticio):
    cte_ficticio.resposta_simulada.write_text('{"linhas": "texto solto"}', encoding="utf-8")
    caso, repositorio = _caso_de_uso(tmp_path, cte_ficticio.pdf.parent)
    saida = tmp_path / "saida.xlsx"

    resultado = caso.executar([cte_ficticio.pdf], saida)

    assert resultado.erros
    assert not saida.exists()
    assert repositorio.execucoes()[-1].resultado == "falhou"


def _alterar_resposta(cte_ficticio, alterar):
    resposta = json.loads(cte_ficticio.resposta_simulada.read_text(encoding="utf-8"))
    alterar(resposta)
    cte_ficticio.resposta_simulada.write_text(json.dumps(resposta), encoding="utf-8")


def _valor(repositorio, resultado, coluna):
    return next(
        v for v in repositorio.valores_do_documento(resultado.documentos[0].id) if v.coluna == coluna
    )


def test_valor_em_branco_conta_como_ausente(tmp_path, cte_ficticio):
    def alterar(resposta):
        resposta["linhas"][0]["campos"]["FORNECEDOR"]["valor"] = "   "

    _alterar_resposta(cte_ficticio, alterar)
    caso, repositorio = _caso_de_uso(tmp_path, cte_ficticio.pdf.parent)

    resultado = caso.executar([cte_ficticio.pdf], tmp_path / "saida.xlsx")

    fornecedor = _valor(repositorio, resultado, "FORNECEDOR")
    assert fornecedor.origem == "ausente"
    assert fornecedor.pendencia == "não encontrado no documento"


def test_pedaco_de_um_numero_nao_conta_como_evidencia(tmp_path, cte_ficticio):
    def alterar(resposta):
        resposta["linhas"][0]["campos"]["VALOR TOTAL"]["valor"] = "1,40"

    _alterar_resposta(cte_ficticio, alterar)
    caso, repositorio = _caso_de_uso(tmp_path, cte_ficticio.pdf.parent)

    resultado = caso.executar([cte_ficticio.pdf], tmp_path / "saida.xlsx")

    assert _valor(repositorio, resultado, "VALOR TOTAL").pendencia == "valor não aparece no trecho citado"


def test_nao_sobrescreve_planilha_existente(tmp_path, cte_ficticio):
    saida = tmp_path / "saida.xlsx"
    saida.write_bytes(b"planilha anterior")
    caso, repositorio = _caso_de_uso(tmp_path, cte_ficticio.pdf.parent)

    resultado = caso.executar([cte_ficticio.pdf], saida)

    assert saida.read_bytes() == b"planilha anterior"
    assert resultado.destino is None
    assert resultado.erros[0].mensagem == "a planilha de destino já existe"
    assert repositorio.execucoes()[-1].resultado == "falhou"


def test_substitui_planilha_existente_quando_pedido(tmp_path, cte_ficticio):
    saida = tmp_path / "saida.xlsx"
    saida.write_bytes(b"planilha anterior")
    caso, _ = _caso_de_uso(tmp_path, cte_ficticio.pdf.parent)

    resultado = caso.executar([cte_ficticio.pdf], saida, substituir=True)

    assert resultado.destino == saida
    assert len(_linhas_da_planilha(saida)[1]) == 1


def test_ia_sem_nenhuma_linha_e_erro(tmp_path, cte_ficticio):
    _alterar_resposta(cte_ficticio, lambda resposta: resposta.update(linhas=[]))
    caso, repositorio = _caso_de_uso(tmp_path, cte_ficticio.pdf.parent)
    saida = tmp_path / "saida.xlsx"

    resultado = caso.executar([cte_ficticio.pdf], saida)

    assert "nenhuma linha" in resultado.erros[0].mensagem
    assert not saida.exists()
    assert repositorio.execucoes()[-1].resultado == "falhou"


class _ExportadorQueFalha:
    def __init__(self, erro):
        self.erro = erro

    def exportar(self, linhas, colunas, destino):
        raise self.erro


def test_planilha_bloqueada_falha_a_execucao_sem_quebrar(tmp_path, cte_ficticio):
    repositorio = RepositorioSqlite(tmp_path / "a.db")
    caso = ProcessarDocumentos(
        LeitorPdf(),
        IASimulada(cte_ficticio.pdf.parent),
        repositorio,
        _ExportadorQueFalha(PermissionError("arquivo aberto no Excel")),
    )

    resultado = caso.executar([cte_ficticio.pdf], tmp_path / "saida.xlsx")

    assert resultado.destino is None
    assert "não foi possível gravar a planilha" in resultado.erros[0].mensagem
    assert repositorio.execucoes()[-1].resultado == "falhou"


def test_erro_inesperado_fica_registrado_no_historico(tmp_path, cte_ficticio):
    repositorio = RepositorioSqlite(tmp_path / "a.db")
    caso = ProcessarDocumentos(
        LeitorPdf(), IASimulada(cte_ficticio.pdf.parent), repositorio, _ExportadorQueFalha(RuntimeError("bug"))
    )

    with pytest.raises(RuntimeError):
        caso.executar([cte_ficticio.pdf], tmp_path / "saida.xlsx")

    ultima = repositorio.execucoes()[-1]
    assert ultima.resultado == "falhou"
    assert ultima.fim is not None


def test_nfse_ficticia_confere_com_o_gabarito(tmp_path, nfse_ficticia):
    caso, _ = _caso_de_uso(tmp_path, nfse_ficticia.pdf.parent)
    saida = tmp_path / "saida.xlsx"

    caso.executar([nfse_ficticia.pdf], saida)

    gabarito = json.loads(nfse_ficticia.gabarito.read_text(encoding="utf-8"))["linhas"][0]
    cabecalho, linhas = _linhas_da_planilha(saida)
    obtida = dict(zip(cabecalho, linhas[0]))
    for coluna in COLUNAS_CONTROLE:
        assert obtida[coluna] == gabarito.get(coluna), coluna
    assert obtida["VALOR TOTAL"] == "1.140,00"  # valor do serviço, não o líquido
    assert obtida["PEDIDO"] == "4500012345"


def test_tipo_de_nota_vem_de_regra_e_nao_da_ia(tmp_path, cte_ficticio, nfse_ficticia):
    def alterar(resposta):
        resposta["linhas"][0]["campos"]["TIPO DE NOTA"] = {
            "valor": "CT-e", "trecho": "DOCUMENTO FICTÍCIO", "pagina": 1, "confianca": 0.9
        }

    _alterar_resposta(nfse_ficticia, alterar)
    caso, repositorio = _caso_de_uso(tmp_path, nfse_ficticia.pdf.parent)

    resultado = caso.executar([cte_ficticio.pdf, nfse_ficticia.pdf], tmp_path / "saida.xlsx")

    tipos = {nome: linha.valores["TIPO DE NOTA"] for nome, linha in resultado.linhas}
    assert tipos["cte_1234.pdf"].valor == "CT-e"
    assert tipos["nfse_5501.pdf"].valor == "Nota Fiscal"
    assert {str(v.origem) for v in tipos.values()} == {"derivado_regra"}
