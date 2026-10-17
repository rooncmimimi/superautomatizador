"""Revisão: ler propõe linhas, o usuário corrige e aprova, e só as aprovadas vão para a planilha."""

import sqlite3

from openpyxl import load_workbook

from automatizador.adaptadores.exportador_xlsx import ExportadorXlsx
from automatizador.adaptadores.ia_simulada import IASimulada
from automatizador.adaptadores.leitor_pdf import LeitorPdf
from automatizador.adaptadores.repositorio_sqlite import RepositorioSqlite
from automatizador.aplicacao.processar_documento import ProcessarDocumentos


def _caso(tmp_path, pasta):
    repositorio = RepositorioSqlite(tmp_path / "a.db")
    return ProcessarDocumentos(LeitorPdf(), IASimulada(pasta), repositorio, ExportadorXlsx()), repositorio


def _planilha(caminho):
    linhas = list(load_workbook(caminho)["Controle"].iter_rows(values_only=True))
    return [dict(zip(linhas[0], linha)) for linha in linhas[1:]]


def test_ler_nao_exporta_e_deixa_a_execucao_em_revisao(tmp_path, cte_ficticio):
    caso, repositorio = _caso(tmp_path, cte_ficticio.pdf.parent)

    lote = caso.ler([cte_ficticio.pdf])

    assert len(lote.linhas) == 1
    assert lote.destino is None
    assert not lote.linhas[0][1].aprovada
    assert repositorio.execucoes()[-1].resultado == "aguardando_revisao"


def test_sem_linha_aprovada_nada_e_gravado(tmp_path, cte_ficticio):
    caso, repositorio = _caso(tmp_path, cte_ficticio.pdf.parent)
    lote = caso.ler([cte_ficticio.pdf])
    saida = tmp_path / "saida.xlsx"

    caso.exportar(lote, saida)

    assert not saida.exists()
    assert lote.erros[-1].mensagem == "nenhuma linha aprovada para exportar"
    assert repositorio.execucoes()[-1].resultado == "falhou"


def test_so_linhas_aprovadas_sao_gravadas(tmp_path, cte_ficticio, nfse_ficticia):
    caso, repositorio = _caso(tmp_path, cte_ficticio.pdf.parent)
    lote = caso.ler([cte_ficticio.pdf, nfse_ficticia.pdf])
    nfse = next(linha for nome, linha in lote.linhas if nome == "nfse_5501.pdf")
    saida = tmp_path / "saida.xlsx"

    caso.aprovar([nfse], "ana")
    caso.exportar(lote, saida)

    linhas = _planilha(saida)
    assert [linha["CT-e"] for linha in linhas] == ["5501"]
    assert repositorio.aprovacao_da_linha(nfse.id)[0] == "ana"
    assert repositorio.execucoes()[-1].resultado == "concluida"


def test_correcao_vira_nova_versao_com_usuario_e_vai_para_a_planilha(tmp_path, cte_ficticio):
    caso, repositorio = _caso(tmp_path, cte_ficticio.pdf.parent)
    lote = caso.ler([cte_ficticio.pdf])
    linha = lote.linhas[0][1]
    caso.aprovar([linha], "ana")

    caso.corrigir(linha, "CÓDIGO FORNECEDOR", " 3100289 ", "ana")

    assert not linha.aprovada  # corrigiu depois de aprovar: precisa aprovar de novo
    assert repositorio.aprovacao_da_linha(linha.id) == (None, None)
    versoes = repositorio.versoes_do_valor(linha.id, "CÓDIGO FORNECEDOR")
    assert [(v.versao, v.valor, v.origem, v.usuario) for v in versoes] == [
        (1, None, "ausente", None),
        (2, "3100289", "manual", "ana"),
    ]
    atual = next(v for v in repositorio.valores_do_documento(lote.documentos[0].id) if v.coluna == "CÓDIGO FORNECEDOR")
    assert atual.valor == "3100289"

    caso.aprovar([linha], "ana")
    caso.exportar(lote, tmp_path / "saida.xlsx")
    assert _planilha(tmp_path / "saida.xlsx")[0]["CÓDIGO FORNECEDOR"] == "3100289"


def test_correcao_fora_do_formato_continua_pendente(tmp_path, cte_ficticio):
    caso, _ = _caso(tmp_path, cte_ficticio.pdf.parent)
    linha = caso.ler([cte_ficticio.pdf]).linhas[0][1]

    valor = caso.corrigir(linha, "VALOR TOTAL", "mil reais", "ana")

    assert valor.pendencia == "valor não está no formato 1.234,56"
    assert linha.pendencias == ["VALOR TOTAL: valor não está no formato 1.234,56"]


def test_apagar_um_valor_obrigatorio_deixa_a_linha_pendente(tmp_path, cte_ficticio):
    caso, _ = _caso(tmp_path, cte_ficticio.pdf.parent)
    linha = caso.ler([cte_ficticio.pdf]).linhas[0][1]

    caso.corrigir(linha, "FORNECEDOR", "", "ana")

    assert linha.pendencias == ["FORNECEDOR: não encontrado no documento"]


def test_banco_de_versao_anterior_e_atualizado(tmp_path):
    banco = tmp_path / "antigo.db"
    conexao = sqlite3.connect(banco)
    conexao.executescript(
        "CREATE TABLE linha (id INTEGER PRIMARY KEY, documento_id INTEGER NOT NULL, ordem INTEGER NOT NULL);"
        "CREATE TABLE valor_campo (id INTEGER PRIMARY KEY, linha_id INTEGER NOT NULL, coluna TEXT NOT NULL,"
        " valor TEXT, origem TEXT NOT NULL, trecho TEXT, pagina INTEGER, confianca REAL, pendencia TEXT,"
        " versao INTEGER NOT NULL DEFAULT 1, criado_em TEXT NOT NULL);"
    )
    conexao.close()

    RepositorioSqlite(banco).fechar()

    conexao = sqlite3.connect(banco)
    assert {"aprovada_por", "aprovada_em"} <= {c[1] for c in conexao.execute("PRAGMA table_info(linha)")}
    assert "usuario" in {c[1] for c in conexao.execute("PRAGMA table_info(valor_campo)")}
    conexao.close()
