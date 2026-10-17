"""Testes de fumaça da linha de comando e da janela."""

from openpyxl import load_workbook

from automatizador.cli import main
from automatizador.configuracao import criar_processador
from automatizador.interface.janela import COR_PENDENTE, JanelaPrincipal


def test_linha_de_comando_gera_ficticio_e_processa(tmp_path):
    pasta = tmp_path / "docs"
    saida = tmp_path / "controle.xlsx"

    assert main(["gerar-ficticios", str(pasta)]) == 0
    codigo = main(
        ["processar", str(pasta / "cte_1234.pdf"), "--saida", str(saida), "--banco", str(tmp_path / "a.db")]
    )

    assert codigo == 0
    assert load_workbook(saida)["Controle"].max_row == 2


def test_linha_de_comando_falha_quando_nada_pode_ser_lido(tmp_path):
    invalido = tmp_path / "quebrado.pdf"
    invalido.write_bytes(b"isto nao e um pdf")

    codigo = main(
        ["processar", str(invalido), "--saida", str(tmp_path / "x.xlsx"), "--banco", str(tmp_path / "a.db")]
    )

    assert codigo == 1
    assert not (tmp_path / "x.xlsx").exists()


def test_janela_mostra_linhas_e_destaca_pendentes(qtbot, tmp_path, cte_ficticio):
    janela = JanelaPrincipal(criar_processador(banco=tmp_path / "a.db"))
    qtbot.addWidget(janela)

    janela.definir_arquivos([cte_ficticio.pdf])
    janela.processar(tmp_path / "controle.xlsx")

    assert janela.tabela.rowCount() == 1
    cabecalhos = [janela.tabela.horizontalHeaderItem(i).text() for i in range(janela.tabela.columnCount())]
    fornecedor = janela.tabela.item(0, cabecalhos.index("FORNECEDOR"))
    codigo = janela.tabela.item(0, cabecalhos.index("CÓDIGO FORNECEDOR"))
    assert fornecedor.text() == "Transportadora Exemplo Ltda"
    assert codigo.text() == ""
    assert codigo.background().color() == COR_PENDENTE
